"""Export Part B documents, slides and figures from completed measured outputs.

Run from the repository root after the complete pipeline:
    python -m part_b.make_deliverables

No metric is inserted by hand. Missing, inconsistent or incomplete required
results fail instead of producing documents with invented performance claims.
"""
from __future__ import annotations

import argparse
from html import escape
import json
import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/part_b_matplotlib")
os.environ.setdefault("XDG_CACHE_HOME", "/tmp/part_b_cache")
Path(os.environ["XDG_CACHE_HOME"]).mkdir(parents=True, exist_ok=True)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parent
SOURCE = "https://github.com/erwinzhang7/SSU-Data-Dash"
COMMIT = "6c5831e12f523d617eccd1418d46a5cf2a0bb8b4"
NAVY, TEAL, ORANGE, GREY = "#142D45", "#14867D", "#D57B38", "#536373"
LABELS = {
    "global_median": "Global median",
    "rolling_community_12m": "Community/type 12-month median",
    "nearest_sales_10": "10 nearest earlier sales",
    "random_forest": "Random forest",
    "hist_gradient_boosting": "Histogram gradient boosting",
    "catboost_d6_rmse": "CatBoost depth 6 / RMSE",
    "catboost_d7_rmse": "CatBoost depth 7 / RMSE",
    "catboost_d7_mae": "CatBoost depth 7 / MAE",
    "catboost_no_remarks": "CatBoost depth 6 / no remarks",
}


def label(name):
    return LABELS.get(name, str(name).replace("_", " "))


def read_results(folder):
    required = ["metrics_summary.json", "selection.json", "validation_predictions.csv", "error_groups.csv",
                "models/model_metadata.json"]
    missing = [str(folder / f) for f in required if not (folder / f).is_file()]
    if missing:
        raise FileNotFoundError("Run the complete pipeline before exporting; missing: " + ", ".join(missing))
    summary = json.loads((folder / "metrics_summary.json").read_text())
    selection = json.loads((folder / "selection.json").read_text())
    if summary["fingerprint"] != selection["fingerprint"]:
        raise ValueError("Summary and selection belong to different pipeline runs.")
    metadata = json.loads((folder / "models" / "model_metadata.json").read_text())
    if metadata["fingerprint"] != summary["fingerprint"] or metadata["train_n"] != summary["training_rows"]:
        raise ValueError("Final-model metadata does not match the validated data and configuration.")
    winner = selection["selected_model"]
    if summary["selected_model"] != winner:
        raise ValueError("Selected model differs between summary and selection.")
    validation = pd.read_csv(folder / "validation_predictions.csv")
    selected = validation.loc[validation.model == winner]
    if sorted(selected.month.unique()) != sorted(summary["validation_months"]):
        raise ValueError("Selected-model predictions do not cover every declared validation month.")
    if selected.ml_num.duplicated().any():
        raise ValueError("Selected-model validation IDs are duplicated.")
    for score in summary["validation_metrics"]:
        rows = validation.loc[validation.model == score["model"]]
        ape = 100 * np.abs(rows.predicted_price - rows.sold_price) / rows.sold_price
        if len(rows) != score["n"] or not np.isclose(np.median(ape), score["mdape"]):
            raise ValueError("Validation CSV and reported metrics disagree for " + score["model"])
    result = {
        "summary": summary, "selection": selection, "validation": validation,
        "groups": pd.read_csv(folder / "error_groups.csv"), "winner": winner,
        "score": next(s for s in summary["validation_metrics"] if s["model"] == winner),
        "baseline": next(s for s in summary["validation_metrics"] if s["model"] == "rolling_community_12m"),
    }
    result["relative_gain"] = 100 * (1 - result["score"]["mdape"] / result["baseline"]["mdape"])
    importance = folder / "feature_importance.csv"
    result["importance"] = pd.read_csv(importance) if importance.is_file() else None
    submission = ROOT.parent / "submissions" / "submission_off_market.csv"
    if not submission.is_file():
        raise FileNotFoundError("Generate and check the off-market submission before exporting documents.")
    submitted = pd.read_csv(submission)
    if list(submitted.columns) != ["ml_num", "predicted_price"] or not submitted.ml_num.is_unique:
        raise ValueError("Submission schema or ID uniqueness is invalid.")
    if len(submitted) != 3974 or not np.isfinite(submitted.predicted_price).all() or (submitted.predicted_price <= 0).any():
        raise ValueError("Submission must contain 3,974 positive finite off-market predictions.")
    result["submission_n"] = len(submitted)
    return result


def write_figures(r, folder):
    folder.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.labelcolor": NAVY,
                         "text.color": NAVY, "axes.spines.top": False, "axes.spines.right": False})
    rows = sorted(r["summary"]["validation_metrics"], key=lambda x: x["mdape"], reverse=True)
    fig, ax = plt.subplots(figsize=(8.8, 3.5))
    y = np.arange(len(rows))
    colors = [TEAL if s["model"] == r["winner"] else "#B8C8D4" for s in rows]
    bars = ax.barh(y, [s["mdape"] for s in rows], color=colors, height=.57)
    ax.set_yticks(y, [label(s["model"]) for s in rows], fontsize=10)
    ax.set_xlabel("Median absolute percentage error (%) — lower is better")
    ax.set_title("Frozen model vs. baselines | Dec 2025–May 2026", loc="left", fontweight="bold", pad=14)
    for bar, row in zip(bars, rows):
        ax.text(bar.get_width() + .25, bar.get_y() + bar.get_height()/2,
                f'{row["mdape"]:.2f}%', va="center", fontsize=10, fontweight="bold")
    ax.set_xlim(0, max(s["mdape"] for s in rows) * 1.19)
    ax.grid(axis="x", alpha=.15)
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(folder / "validation_comparison.png", dpi=190, bbox_inches="tight")
    plt.close(fig)

    monthly = pd.DataFrame(r["summary"]["monthly_metrics"])
    fig, ax = plt.subplots(figsize=(8.8, 3.4))
    for name, color, marker in [("rolling_community_12m", GREY, "o"), (r["winner"], TEAL, "o")]:
        frame = monthly[monthly.model == name].sort_values("month")
        ax.plot(frame.month, frame.mdape, color=color, marker=marker, linewidth=2,
                label=label(name))
    ax.set_ylabel("MdAPE (%)")
    ax.set_title("Monthly walk-forward error", loc="left", fontweight="bold", pad=12)
    ax.set_ylim(bottom=0)
    ax.grid(axis="y", alpha=.15)
    ax.legend(frameon=False, fontsize=9)
    fig.tight_layout()
    fig.savefig(folder / "monthly_walk_forward.png", dpi=190, bbox_inches="tight")
    plt.close(fig)

    bands = r["groups"].loc[r["groups"].group_by == "price_band"].copy()
    band_order = ["Under $500k", "$500k-$1m", "$1m-$2m", "$2m+"]
    bands["order"] = bands.group.map({v: i for i, v in enumerate(band_order)})
    bands = bands.sort_values("order")
    fig, ax = plt.subplots(figsize=(8, 3.2))
    ax.bar(bands.group, bands.mdape, color=TEAL, width=.58)
    for i, row in enumerate(bands.itertuples()):
        ax.text(i, row.mdape + .25, f"{row.mdape:.1f}%\nn={row.n:,}", ha="center", fontsize=9)
    ax.set_ylim(0, max(bands.mdape) * 1.35)
    ax.set_ylabel("MdAPE (%)")
    ax.set_title("Selected model: error by actual sale-price band", loc="left", fontweight="bold")
    ax.grid(axis="y", alpha=.15)
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(folder / "error_by_price_band.png", dpi=190, bbox_inches="tight")
    plt.close(fig)

    if r["importance"] is not None:
        top = r["importance"].nlargest(12, "importance").sort_values("importance")
        fig, ax = plt.subplots(figsize=(8, 4))
        ax.barh(top.feature.str.replace("_", " "), top.importance, color=TEAL)
        ax.set_xlabel("CatBoost feature importance (relative; not a causal effect)")
        ax.set_title("What the final model uses", loc="left", fontweight="bold")
        fig.tight_layout()
        fig.savefig(folder / "feature_importance.png", dpi=190, bbox_inches="tight")
        plt.close(fig)


def table_md(rows, tuning=False):
    lines = ["| Model | Homes | MdAPE | MAPE |", "|---|---:|---:|---:|"]
    for row in rows:
        lines.append(f'| {label(row["model"])} | {row["n"]:,} | {row["mdape"]:.2f}% | {row["mape"]:.2f}% |')
    return "\n".join(lines)


def describe_config(r):
    config = dict(r["selection"]["selected_config"])
    kind = config.pop("kind")
    text = config.pop("include_text", True)
    values = "; ".join(f"{k}={v}" for k, v in config.items())
    return f"{values}; remarks features={'on' if text else 'off'}; log-price target; random seed 42."


def method_text(r):
    family = r["selection"]["selected_config"]["kind"]
    if family == "cat":
        algorithm = "CatBoost learns nonlinear interactions and handles location/type categories and missing numeric values natively."
    elif family == "rf":
        algorithm = "The random forest averages randomized decision trees; category encoding and numeric imputation are fitted within each training fold."
    else:
        algorithm = "Histogram gradient boosting builds trees sequentially; category encoding and numeric imputation are fitted within each training fold."
    text = r["selection"]["selected_config"].get("include_text", True)
    features = ("size and room counts, property type/style, synthetic location, lot dimensions with missingness indicators, "
                "and listing-date calendar features")
    if text:
        features += ", plus fixed remarks flags for condition, amenities and layout"
    return (f"Predict log(sold_price), then exponentiate to dollars. Inputs cover {features}. "
            "Area ranges are represented by their midpoint and width. " + algorithm)


def narrative(r):
    s, b = r["score"], r["baseline"]
    gain = r["relative_gain"]
    comparison = (f"This is {gain:.1f}% lower MdAPE than the rolling community/type baseline "
                  if gain >= 0 else f"This is {-gain:.1f}% higher MdAPE than the rolling community/type baseline ")
    comparison += f"({b['mdape']:.2f}%)."
    expensive = r['groups'].loc[(r['groups'].group_by == 'price_band') & (r['groups'].group == '$2m+')]
    expensive_note = ''
    if len(expensive) == 1:
        row = expensive.iloc[0]
        expensive_note = (f"The $2m+ group had {row.mdape:.2f}% MdAPE across {int(row.n):,} homes, "
                          f"with {row.median_signed_error:+.2f}% median signed error. ")
    return {
        "result": f"{label(r['winner'])} achieved {s['mdape']:.2f}% MdAPE and {s['mape']:.2f}% MAPE over {s['n']:,} held-out sales. {comparison}",
        "protocol": ("Tune on October–November 2025 only, selecting the lowest pooled MdAPE (MAPE breaks ties). "
            "Freeze the model configuration before December 2025–May 2026. For each of those six months, "
            "refit on sales strictly before the month's start and predict that month's held-out sales. "
            "This follows the official sale-month backtest; it is not a strict per-home listing-time simulation."),
        "boundary": ("No list_price, sold_price, date_sold, days_on_market, external_avm_estimate, "
            "external_avm_estimate_date or ml_num enters the learned feature matrix. sold_price is only the target; "
            "date_sold only defines training cutoffs and evaluation groups. No external data is used."),
        "metrics": ("MdAPE = median(100 × |prediction − sale| / sale); MAPE uses the mean. "
            "Both are percentages of actual sale price, and lower is better. MdAPE is the competition's primary metric."),
        "limits": (expensive_note + "The data is synthetic; results do not establish real-world appraisal accuracy. "
            "Six months do not capture every future market regime. Expensive or unusual homes may have larger errors. "
            "The June–September 2026 test prices are hidden, so test accuracy is unknown. "
            "Do not describe the validation score as a hidden-test score."),
    }


def write_text(r, folder):
    n = narrative(r)
    s = r["score"]
    config = r["selection"]["selected_config"]
    key_parameters = {
        "cat": ["depth", "iterations", "learning_rate", "loss_function"],
        "rf": ["n_estimators", "max_depth", "min_samples_leaf", "max_features"],
        "hgb": ["max_iter", "learning_rate", "max_leaf_nodes", "l2_regularization"],
    }[config["kind"]]
    short_config = "; ".join(f"{key}={config[key]}" for key in key_parameters)
    feature_sentence = (
        "Inputs describe area, rooms, parking, type/style, synthetic location, lot dimensions and listing-date calendar features. "
        "Area ranges retain midpoint/width; lot missingness is preserved."
    )
    if config.get("include_text", True):
        feature_sentence += " Fixed remarks flags describe condition, amenities and layout."
    text = ["# Part B — Off-Market Prediction", "", n["result"], "",
        "**Model.** Predict log(sold_price), then exponentiate to dollars. " + feature_sentence, "",
        "Predictors exclude list price, sale-time fields, AVM estimates and IDs. "
        "Price supplies labels; sale date defines folds. No external data.", "",
        "Log-price fitting emphasizes relative error. Categorical processing and any numeric imputation are fitted "
        "using only each fold's training rows; feature engineering is deterministic and row-local.", "",
        "**Selection and evaluation.** Compare baselines, random forest, histogram gradient boosting and CatBoost. "
        "Tune only on October–November 2025 MdAPE; MAPE breaks ties. "
        "Freeze settings before December 2025–May 2026 validation. "
        "Each month, train strictly before month-start and predict that month's sales. "
        "This official sale-month protocol is not a strict per-home listing-time simulation.", "",
        table_md(r["summary"]["validation_metrics"]), "",
        "Rows share the same untouched six-month validation homes. "
        "MdAPE is median absolute percentage error relative to sale price; MAPE is its mean. Lower is better.", "",
        f"**Frozen settings:** `{short_config}`; random seed 42. "
        "Full settings/tuning results: `model_comparison.md`.", "",
        f"**Interpretation.** {s['within_20_pct']:.1f}% of held-out predictions are within 20% of sale price. "
        "Refit the selected model on all 150,000 training sales to generate the off-market submission. "
        "The CSV contains all 3,974 supplied off-market IDs once.", "",
        "**Limits.** Synthetic data cannot establish real-world appraisal accuracy; unusual homes and market shifts remain risks. "
        "June–September 2026 test prices are hidden: these validation results do not measure hidden-test accuracy.", "",
        f"Source: official handout/starter, {SOURCE}, commit `{COMMIT}`.", ""]
    (folder / "report_section.md").write_text("\n".join(text), encoding="utf-8")

    speaker = (f"My part estimates a home's sale price before we know its listing price. "
        f"I compared simple price baselines, a random forest, histogram gradient boosting, and several CatBoost configurations. "
        f"I tuned only on October and November 2025, then froze {label(r['winner'])} before evaluating December through May.\n\n"
        f"For each validation month, the model learned only from sales before that month. "
        f"Its inputs describe size, rooms, property type, synthetic location and listing date. "
        + ("Fixed description keywords add condition and amenity information. "
           if r["selection"]["selected_config"].get("include_text", True) else "The selected version omits description keywords. ")
        + f"List price, post-sale fields and external valuations are excluded.\n\n"
        f"Across {s['n']:,} held-out homes, median absolute percentage error was {s['mdape']:.2f} percent, "
        f"and mean absolute percentage error was {s['mape']:.2f} percent. "
        f"The community baseline's median error was {r['baseline']['mdape']:.2f} percent.\n\n"
        "These are measured validation results, not hidden-test scores. "
        "Synthetic data and market changes limit generalization. "
        "The final deliverable is a reproducible model and one prediction for every off-market test home.")
    words = len(speaker.split())
    (folder / "speaker_script_75_seconds.md").write_text(
        f"# Part B speaker script\n\nApprox. {words} words; rehearse at about {words/1.25:.0f} words/minute for 75 seconds.\n\n"
        + speaker + "\n", encoding="utf-8")

    chinese = f"""# B 部分交接说明：无挂牌价预测

你的任务是：在不知道挂牌价时，预测房屋最终成交价。这里已完成模型比较、调参、按时间验证、最终模型和测试集预测；不包含 C 部分的有挂牌价模型。

最终模型：**{label(r['winner'])}**。先仅用 2025 年 10–11 月结果选模型，再固定参数，用 2025 年 12 月至 2026 年 5 月的六个月做逐月向前验证。每月只训练该月开始以前已成交的数据，不随机划分。

实际验证结果：**MdAPE {s['mdape']:.2f}%、MAPE {s['mape']:.2f}%**，共 **{s['n']:,}** 套验证房屋。社区/房型过去一年中位数基线的 MdAPE 为 **{r['baseline']['mdape']:.2f}%**。这些是代码真实计算的验证分数，不能说成隐藏测试集分数。

MdAPE 是每套房“预测与真实成交价之差的绝对值÷真实成交价”的中位数；MAPE 是这些百分比误差的平均数。两者越小越好。

特征包含面积、卧室/卫浴、停车、房型、合成地理位置、挂牌日期，以及所选模型启用时的描述关键词。面积区间取中点，同时保留区间宽度；地块缺失值保留缺失信息。学习模型不使用挂牌价、成交日期、市场停留天数、外部估值或 ID。成交价只作为训练目标，成交日期只用于时间切分。

注意：官方验证按“成交月份”切分，不能声称这是严格在每套房挂牌当天完成的历史估值。数据是合成数据，不能直接承诺真实房地产市场效果。隐藏测试成交价未提供，所以不能计算测试准确率。

交给团队：`report_section.md` 可插入全队最多四页的报告；两页独立报告用于审阅，不要直接占用整队四页额度。`speaker_script_75_seconds.md` 为英文讲稿，需自行计时练习；PPT 为 B 部分两张幻灯片。全队仍需自行合并 C 部分文件、填写成员姓名及仓库/视频链接，最后提交。

规则来源：{SOURCE}，版本 `{COMMIT}`。没有使用外部数据，也没有尝试恢复隐藏测试标签。
"""
    (folder / "chinese_handoff.md").write_text(chinese, encoding="utf-8")
    appendix = ["# Measured model-selection results", "",
        "October–November 2025 tuning only. These scores select the model; they are not fresh final validation scores.", "",
        table_md(r["selection"]["tuning_metrics"]), "", "Selection rule: " + r["selection"]["selection_rule"], "",
        "## Frozen configuration", "", "```json", json.dumps(r["selection"]["selected_config"], indent=2), "```", "",
        "## Untouched final walk-forward backtest", "", table_md(r["summary"]["validation_metrics"]), ""]
    (folder / "model_comparison.md").write_text("\n".join(appendix), encoding="utf-8")


def write_pdf(r, folder, figures):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak

    font_root = Path("/usr/share/fonts/truetype/dejavu")
    if (font_root / "DejaVuSans.ttf").is_file():
        pdfmetrics.registerFont(TTFont("PartB", str(font_root / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("PartBBold", str(font_root / "DejaVuSans-Bold.ttf")))
        body_font, bold_font = "PartB", "PartBBold"
    else:
        body_font, bold_font = "Helvetica", "Helvetica-Bold"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("BBody", fontName=body_font, fontSize=9, leading=12.7, textColor=colors.HexColor(NAVY), spaceAfter=7))
    styles.add(ParagraphStyle("BTitle", fontName=bold_font, fontSize=20, leading=25, textColor=colors.HexColor(NAVY), spaceAfter=8))
    styles.add(ParagraphStyle("BHead", fontName=bold_font, fontSize=11, leading=14, textColor=colors.HexColor(TEAL), spaceBefore=7, spaceAfter=5))
    styles.add(ParagraphStyle("BSmall", fontName=body_font, fontSize=7.5, leading=10, textColor=colors.HexColor(GREY), spaceAfter=6))
    styles.add(ParagraphStyle("BCell", fontName=body_font, fontSize=8, leading=10, textColor=colors.HexColor(NAVY)))
    n = narrative(r)
    p = lambda text, style="BBody": Paragraph(escape(str(text)), styles[style])
    def score_table(rows):
        data = [["Model", "MdAPE", "MAPE", "Homes"]]
        data += [[p(label(s["model"]), "BCell"), f'{s["mdape"]:.2f}%', f'{s["mape"]:.2f}%', f'{s["n"]:,}'] for s in rows]
        table = Table(data, colWidths=[255, 65, 65, 65], hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), bold_font),
            ("FONTNAME", (0, 1), (-1, -1), body_font),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#EDF4F7"), colors.white]),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ]))
        return table
    story = [p("Part B — Off-Market Prediction", "BTitle"),
        p("SSU Data Dash 2026 | Model development and measured results", "BSmall"),
        p(n["result"]), p("Model and information boundary", "BHead"),
        p(method_text(r)), p(n["boundary"]),
        p("Untouched six-month walk-forward backtest", "BHead"), p(n["protocol"]),
        score_table(r["summary"]["validation_metrics"]), Spacer(1, 8),
        Image(str(figures / "monthly_walk_forward.png"), width=450, height=174),
        p(n["metrics"], "BSmall"), PageBreak(),
        p("Selection, final model and limits", "BTitle"),
        p("Earlier tuning: October–November 2025", "BHead"),
        p("The table below is used only for model selection. It is deliberately separated from the later six-month backtest.", "BSmall"),
        score_table(r["selection"]["tuning_metrics"]), Spacer(1, 7),
        p("Frozen selected configuration", "BHead"), p(describe_config(r), "BSmall"),
        p("Interpretation and deliverable", "BHead"),
        p(f"The selected model puts {r['score']['within_10_pct']:.1f}% of held-out predictions within 10% of sale price "
          f"and {r['score']['within_20_pct']:.1f}% within 20%. Median signed error: {r['score']['median_signed_error']:+.2f}%. "
          "The final model is refitted on all 150,000 training sales, then predicts the supplied off-market test homes."),
        p(n["limits"]),
        p("Report use and provenance", "BHead"),
        p("This standalone Part B handoff is not the complete team submission. Insert report_section.md into the team's "
          "four-page report and add member names, GitHub and video links. The Part B video allocation is approximately 75 seconds.", "BSmall"),
        p(f"Official handout/starter: {SOURCE}. Reference commit: {COMMIT}. No external data was used.", "BSmall")]
    def footer(canvas, doc):
        canvas.setFont(body_font, 7)
        canvas.setFillColor(colors.HexColor(GREY))
        canvas.drawString(43, 29, "Part B | Off-market | Validation scores only; hidden-test accuracy unknown")
        canvas.drawRightString(569, 29, str(doc.page))
    doc = SimpleDocTemplate(str(folder / "part_b_report.pdf"), pagesize=(612, 792),
        leftMargin=43, rightMargin=43, topMargin=37, bottomMargin=43,
        title="Part B — Off-Market Prediction", author="SSU Data Dash team")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def write_docx(r, folder, figures):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.55)
    section.left_margin = section.right_margin = Inches(.65)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.04
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor.from_string(NAVY[1:])
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 2"].font.size = Pt(11)
    n = narrative(r)
    def table(rows):
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Shading Accent 1"
        for cell, text in zip(t.rows[0].cells, ["Model", "MdAPE", "MAPE", "Homes"]):
            cell.text = text
        for s in rows:
            for cell, text in zip(t.add_row().cells, [label(s["model"]), f'{s["mdape"]:.2f}%', f'{s["mape"]:.2f}%', f'{s["n"]:,}']):
                cell.text = text
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    for run in para.runs:
                        run.font.size = Pt(9)
    doc.add_heading("Part B — Off-Market Prediction", 0)
    doc.add_paragraph("SSU Data Dash 2026 | Model development and measured results", "Subtitle")
    doc.add_paragraph(n["result"])
    doc.add_heading("Model and information boundary", 1)
    doc.add_paragraph(method_text(r))
    doc.add_paragraph(n["boundary"])
    doc.add_heading("Untouched six-month walk-forward backtest", 1)
    doc.add_paragraph(n["protocol"])
    table(r["summary"]["validation_metrics"])
    doc.add_picture(str(figures / "monthly_walk_forward.png"), width=Inches(6.65))
    p = doc.add_paragraph(n["metrics"])
    for run in p.runs:
        run.font.size = Pt(8)
    doc.add_page_break()
    doc.add_heading("Selection, final model and limits", 0)
    doc.add_heading("Earlier tuning: October–November 2025", 1)
    doc.add_paragraph("These earlier months select the model; they are separate from the final six-month backtest.")
    table(r["selection"]["tuning_metrics"])
    doc.add_heading("Frozen selected configuration", 1)
    doc.add_paragraph(describe_config(r))
    doc.add_heading("Interpretation and deliverable", 1)
    doc.add_paragraph(f"{r['score']['within_10_pct']:.1f}% of held-out predictions are within 10% of sale price; "
        f"{r['score']['within_20_pct']:.1f}% are within 20%. Median signed error: {r['score']['median_signed_error']:+.2f}%. "
        "Refit the final selected model on all 150,000 training sales and predict each supplied off-market test ID.")
    doc.add_paragraph(n["limits"])
    doc.add_heading("Report use and provenance", 1)
    doc.add_paragraph("Use report_section.md as the Part B insert in the team's four-page report. "
        "The final team report still needs member names, GitHub and video links; the Part B script targets approximately 75 seconds.")
    doc.add_paragraph(f"Source: {SOURCE}; official commit {COMMIT}. No external data used.")
    section.footer.paragraphs[0].text = "Part B | Validation scores only; hidden-test accuracy unknown"
    doc.save(folder / "part_b_report.docx")


def write_slides(r, folder, figures):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def text(slide, value, x, y, w, h, size=20, color=NAVY, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.text_frame.word_wrap = True
        for i, line in enumerate(value.split("\n")):
            p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold = "Aptos", Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color[1:])
            p.space_after = Pt(12)
        return box
    def slide(title, subtitle):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(250, 252, 253)
        text(s, "PART B  /  OFF-MARKET", .55, .22, 12, .35, 12, TEAL, True)
        text(s, title, .55, .75, 12.2, .65, 30, NAVY, True)
        text(s, subtitle, .58, 1.47, 12, .55, 15, GREY)
        text(s, "SSU Data Dash 2026 | Synthetic data | Hidden-test prices unavailable", .58, 7.1, 12, .25, 10, GREY)
        return s
    s = slide("Estimate sale price before the list price is known", "A reproducible model, selected before the final six-month backtest")
    text(s, "01  FEATURES", .65, 2.15, 3.6, .45, 16, TEAL, True)
    text(s, "Size and room counts\nType, style and synthetic location\nLot size and missingness\nListing-date calendar features" +
         ("\nFixed remarks / condition flags" if r["selection"]["selected_config"].get("include_text", True) else "\nRemarks omitted in selected model"),
         .65, 2.75, 3.7, 3.0, 18)
    text(s, "02  MODEL", 4.85, 2.15, 3.55, .45, 16, TEAL, True)
    text(s, label(r["winner"]) + "\nLearn log(price) → exponentiate\nCompare RF, gradient boosting and CatBoost\nTune Oct–Nov 2025 only",
         4.85, 2.75, 3.6, 3.0, 18)
    text(s, "03  HONEST CHECK", 9.0, 2.15, 3.7, .45, 16, TEAL, True)
    text(s, "Freeze before Dec–May\nRefit using earlier sales each month\nNo list price / post-sale / AVM inputs\nSale-month protocol; not per-home listing-time",
         9.0, 2.75, 3.7, 3.0, 18)
    s.notes_slide.notes_text_frame.text = "Explain the problem, allowed features, log-price model, and separation of tuning from final validation. " + describe_config(r)
    sc = r["score"]
    s = slide("Measured error on future held-out months", f"December 2025–May 2026 | {sc['n']:,} held-out homes | selected configuration unchanged")
    text(s, f"{sc['mdape']:.2f}%", .7, 2.2, 3.1, .9, 42, TEAL, True)
    text(s, "Median absolute % error", .72, 3.12, 3.3, .5, 18)
    text(s, f"{sc['mape']:.2f}% MAPE\n{r['baseline']['mdape']:.2f}% baseline MdAPE\n{sc['within_20_pct']:.1f}% within 20% of sale price", .72, 3.85, 3.6, 1.8, 19)
    s.shapes.add_picture(str(figures / "validation_comparison.png"), Inches(4.3), Inches(2.2), width=Inches(8.45))
    text(s, "Validation ≠ hidden-test accuracy. Synthetic data and future market shifts limit generalization.",
         .72, 6.3, 12, .6, 17, GREY)
    s.notes_slide.notes_text_frame.text = narrative(r)["result"] + " " + narrative(r)["limits"]
    prs.save(folder / "part_b_slides.pptx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs")
    parser.add_argument("--deliverables-dir", type=Path, default=ROOT / "deliverables")
    args = parser.parse_args()
    r = read_results(args.output_dir)
    args.deliverables_dir.mkdir(parents=True, exist_ok=True)
    figures = ROOT / "figures"
    write_figures(r, figures)
    write_text(r, args.deliverables_dir)
    write_pdf(r, args.deliverables_dir, figures)
    write_docx(r, args.deliverables_dir, figures)
    write_slides(r, args.deliverables_dir, figures)
    manifest = {"selected_model": r["winner"], "fingerprint": r["summary"]["fingerprint"],
                "source": SOURCE, "source_commit": COMMIT,
                "files": sorted(p.name for p in args.deliverables_dir.iterdir() if p.is_file())}
    (args.deliverables_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(f"Exported measured Part B documents to {args.deliverables_dir}")


if __name__ == "__main__":
    main()
